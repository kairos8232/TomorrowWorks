"""Convert a DOSM Job Market Insights (JMI) data file into a JSON snapshot.

Usage:
    python3 scripts/import_dosm_jmi.py data/raw/Data_JMI_Q2_2026.xlsx 2026-08-18 2026-10-21

Arguments: the DOSM file, the release date shown on the DOSM release page
(publishedAt), and the date you downloaded the file (retrievedAt).
Output: data/sources/dosm-jmi-<latest quarter>.json

The DOSM file has one sheet with two tables:
  Table 1: job offerings advertised online by economic activity
  Table 2: job offerings advertised online by major occupation group
Each table has annual columns, then quarterly columns (2019 Q3 onward),
then monthly columns. We keep only the quarterly series.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

YEAR_ROW, PERIOD_ROW = 3, 4  # header rows in the DOSM sheet

# ISCO-08 / MASCO 2020 major group codes for Table 2 labels.
MAJOR_GROUP_CODES = {
    "Managers": "1",
    "Professionals": "2",
    "Technicians and associate professionals": "3",
    "Clerical support workers": "4",
    "Service and sales workers": "5",
    "Skilled agricultural, forestry, livestock and fishery workers": "6",
    "Craft and related trades workers": "7",
    "Plant and machine operators and assemblers": "8",
    "Elementary Occupations": "9",
}


def quarter_columns(ws):
    """Return [(column index, 'YYYY-Qn')] for the quarterly block."""
    columns, year = [], None
    for col in range(1, ws.max_column + 1):
        y = ws.cell(YEAR_ROW, col).value
        p = ws.cell(PERIOD_ROW, col).value
        if isinstance(y, int):
            year = y
        if isinstance(p, str) and re.fullmatch(r"Q[1-4]", p) and year:
            columns.append((col, f"{year}-{p}"))
    return columns


def read_table(ws, title_prefix, columns):
    """Read rows under the table whose title starts with title_prefix."""
    start = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 2).value or "").startswith(title_prefix))
    rows = []
    for r in range(start + 1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if isinstance(label, str) and label.startswith("Source:"):
            break
        if not isinstance(label, str) or label in ("Occupation", "Economic Activity"):
            continue
        series = {}
        for col, quarter in columns:
            value = ws.cell(r, col).value
            series[quarter] = value if isinstance(value, (int, float)) else None
        rows.append({"label": "Total" if label == "Jumlah" else label.strip(), "quarterly": series})
    return rows


def main(xlsx_path, published_at, retrieved_at):
    ws = openpyxl.load_workbook(xlsx_path, data_only=True).active
    columns = quarter_columns(ws)
    if not columns:
        sys.exit("No quarterly columns found: the DOSM layout may have changed.")

    occupations = read_table(ws, "Table 2", columns)
    activities = read_table(ws, "Table 1", columns)

    # The latest quarter is the last one with a Total value.
    total = next(r for r in occupations if r["label"] == "Total")
    latest = [q for q, v in total["quarterly"].items() if v is not None][-1]

    for row in occupations:
        if row["label"] != "Total":
            code = MAJOR_GROUP_CODES.get(row["label"])
            if code is None and row["label"] != "Unclassified":
                sys.exit(f"Unknown occupation group: {row['label']!r}")
            row["iscoMajorGroup"] = code

    year, q = latest.split("-")
    snapshot = {
        "source": {
            "name": f"DOSM Big Data Analytics: Job Market Insights, {q} {year}",
            "publisher": "Department of Statistics Malaysia (DOSM)",
            "url": f"https://www.dosm.gov.my/portal-main/release-content/big-data-analytics-job-market-insights-{q.lower()}{year}",
            "file": Path(xlsx_path).name,
            "latestQuarter": latest,
            "publishedAt": published_at,
            "retrievedAt": retrieved_at,
            "geography": "Malaysia",
            "measure": "Number of job offerings advertised online on private recruitment websites",
            "coverage": "Online job advertisements only; not all job vacancies in Malaysia. "
                        "Occupation detail is by major group (1-digit MASCO).",
        },
        "occupationMajorGroups": occupations,
        "economicActivities": activities,
    }
    out = Path("data/sources") / f"dosm-jmi-{latest.lower()}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {out} (latest quarter {latest}, {len(occupations)} occupation rows, "
          f"{len(activities)} activity rows)")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3])
